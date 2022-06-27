from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
import math
import re
import pydash

from helpers.utils import check_tag_exist, replace_tags, get_tag_content, get_tag_from_string, eval_executor
from helpers.texts import text_tag_update
from helpers.images import replace_images
from helpers.tables import replace_tables


def _if(presentation, slide, shape,slides_index, replacements):
    pattern = r'\+\+\+IF (.*?)IF-END\+\+\+'
    matches = get_tag_content(pattern, shape)
    
    if( not matches or len(matches) < 1):
        return

    for match in matches:
        pattern_condition = r'\(\((.*?)\)\)'
        matched_condition = get_tag_from_string(pattern_condition,match)

        pattern_content = r'\<\<(.*?)\>\>'
        matched_content = get_tag_from_string(pattern_content,match)
        
        for contidion in matched_condition:
            object_value = eval_executor(contidion, replacements)

            #replace text
            text_replace_pattern = r'\+\+\+INS (.*?) \+\+\+'
            text_matches = get_tag_from_string(text_replace_pattern, matched_content[0])

            if( text_matches and len(text_matches) > 0):
                text_replace = ""
                if(object_value):
                    updated_data = text_tag_update(matched_content[0],replacements)
                    if(updated_data and updated_data["text"]):
                        text_replace = updated_data["text"]
                    
                # this is not working if you use tabspaces, but you can use spaces
                replace_tags(str(f"+++IF {match}IF-END+++"), text_replace, shape)


            #remove tables
            table_remove_pattern = r'\+\+\+TABLE_REMOVE (.*?) \+\+\+'
            table_remove_matches = get_tag_from_string(table_remove_pattern, matched_content[0])
            
            if( table_remove_matches and len(table_remove_matches) > 0):
                if(object_value):
                    table_remove_index_matches = table_remove_matches[0]
                    _shap_count = 0
                    for _shape in slide.shapes:
                        if _shape.has_table: 
                            for row in _shape.table.rows:
                                for cell in row.cells:
                                    if table_remove_index_matches in cell.text:
                                        old_picture = slide.shapes[_shap_count]
                                        old_pic = old_picture._element
                                        old_pic.getparent().remove(old_pic)
                                        replace_tags(str(f"+++IF {match}IF-END+++"), "", shape)
                                        break
                        _shap_count += 1
                    
        

